"""Read-only input adapters and exact excerpt linking."""
from collections import Counter, defaultdict
import json
from pathlib import Path
import re

from .models import Comment, digest


def _unique_object(pairs):
    obj = {}
    for key, value in pairs:
        if key in obj:
            raise ValueError(f"Duplicate JSON key: {key}")
        obj[key] = value
    return obj


def load_comments(path):
    path = Path(path)
    data = path.read_bytes()
    rows = json.loads(data.decode("utf-8-sig"), object_pairs_hook=_unique_object)
    if not isinstance(rows, list):
        raise ValueError("JSON root must be an array")
    file_hash = digest(data)
    comments = []
    for i, row in enumerate(rows):
        if not isinstance(row, dict) or set(row) != {"교수", "강의명", "댓글"}:
            raise ValueError(f"Unexpected input schema at row {i}")
        if any(not isinstance(row[k], str) or not row[k].strip() for k in ("교수", "강의명")):
            raise ValueError(f"Missing instructor/title at row {i}")
        if not isinstance(row["댓글"], list):
            raise ValueError(f"Comments must be an array at row {i}")
        for j, text in enumerate(row["댓글"]):
            if not isinstance(text, str):
                raise ValueError(f"Comment must be a string at /{i}/댓글/{j}")
            comments.append(Comment(i, j, row["강의명"], row["교수"], text, path.name, file_hash))
    norm = lambda value: re.sub(r"\s+", "", value).lower()
    counts = Counter(c.text for c in comments)
    profile = {
        "file": path.name, "sha256": file_hash,
        "course_rows": len(rows),
        "unique_course_instructor_pairs": len({(r["강의명"], r["교수"]) for r in rows}),
        "unique_titles": len({r["강의명"] for r in rows}),
        "unique_instructor_labels": len({r["교수"] for r in rows}),
        "normalized_unique_titles": len({norm(r["강의명"]) for r in rows}),
        "normalized_unique_instructor_labels": len({norm(r["교수"]) for r in rows}),
        "normalized_unique_pairs": len({(norm(r["강의명"]), norm(r["교수"])) for r in rows}),
        "comments": len(comments), "empty_comments": sum(not c.text.strip() for c in comments),
        "rows_without_comments": sum(not r["댓글"] for r in rows),
        "unique_comment_texts": len(counts),
        "duplicate_text_occurrences": sum(n - 1 for n in counts.values()),
    }
    return comments, profile


def link_workbook(path, comments):
    # Import lazily: synthetic extractor tests need neither openpyxl nor private inputs.
    from openpyxl import load_workbook

    path = Path(path)
    before = digest(path.read_bytes())
    index = defaultdict(list)
    for comment in comments:
        index[(comment.instructor, comment.title)].append(comment)
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    links, sheets = [], []
    pairs, titles, instructors = set(), set(), set()
    try:
        for sheet in workbook:
            rows = sheet.iter_rows()
            header = next(rows, None)
            if header is None or [c.value for c in header] != ["교수", "강의", "통계량"]:
                raise ValueError(f"Unsupported XLSX header in {sheet.title}")
            count = 0
            for cells in rows:
                values = [c.value for c in cells]
                if all(v is None for v in values):
                    continue
                if len(values) != 3 or any(c.data_type == "f" for c in cells):
                    raise ValueError("XLSX must contain three text columns, without formulas")
                if any(not isinstance(v, str) or not v.strip() for v in values):
                    raise ValueError("Blank or non-text XLSX entry")
                professor, title, excerpt = values
                pairs.add((professor, title)); titles.add(title); instructors.add(professor)
                matches = []
                for comment in index.get((professor, title), []):
                    start = comment.text.find(excerpt)
                    while start >= 0:
                        end = start + len(excerpt)
                        suffix = comment.text[end:]
                        truncated = bool(
                            re.search(r"\d\.$", excerpt) and re.match(r"\d", suffix)
                            or re.search(r"\d$", excerpt) and re.match(r"\.\d|\d", suffix)
                        )
                        matches.append({
                            "json_pointer": comment.pointer, "start": start, "end": end,
                            "whole_comment": excerpt == comment.text,
                            "truncated_number_at_end": truncated,
                        })
                        start = comment.text.find(excerpt, start + 1)
                links.append({
                    "sheet": sheet.title, "row": cells[0].row,
                    "cell": cells[2].coordinate, "xlsx_sha256": before,
                    "excerpt": excerpt, "matches": matches,
                    "status": "matched" if len(matches) == 1 else "ambiguous" if matches else "unmatched",
                })
                count += 1
            sheets.append({"name": sheet.title, "dimension": sheet.calculate_dimension(), "data_rows": count})
    finally:
        workbook.close()
    if digest(path.read_bytes()) != before:
        raise RuntimeError("Input XLSX changed during read")
    summary = {
        "file": path.name, "sha256": before, "sheets": sheets, "rows": len(links),
        "unique_course_instructor_pairs": len(pairs), "unique_titles": len(titles),
        "unique_instructor_labels": len(instructors),
        "statuses": dict(Counter(r["status"] for r in links)),
        "whole_comment_rows": sum(len(r["matches"]) == 1 and r["matches"][0]["whole_comment"] for r in links),
        "distinct_matched_comments": len({m["json_pointer"] for r in links for m in r["matches"]}),
        "truncated_number_rows": sum(any(m["truncated_number_at_end"] for m in r["matches"]) for r in links),
    }
    return links, summary
